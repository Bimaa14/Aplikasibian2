"""Replace the active inventory offline; retain old IDs for historical references."""
import asyncio
import hashlib
import json
import re
import sys
import uuid
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from bson import json_util

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from lib.db import db
from lib.financial_operations import FinancialOperation, financial_lock
from lib.runtime_lock import single_writer
from lib.workbooks import dataset
from models.product import Product
from tools.prepare_store_candidate import snapshot, fingerprint

ROOT = Path(__file__).resolve().parents[2]


def normalized(value):
    return re.sub(r'\s+', ' ', str(value or '').strip()).casefold()


def unique_match(row, records):
    candidates = [p for p in records if normalized(p.get('sku')) == normalized(row['sku'])]
    if not candidates and normalized(row['name']):
        candidates = [p for p in records if normalized(p['name']) == normalized(row['name'])]
    if len(candidates) > 1:
        raise ValueError(f"Ambiguous product: {row['sku']} / {row['name']}")
    return candidates[0] if candidates else None


def read_inventory(path):
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook['Laporan Data Barang']
        headers = list(next(sheet.iter_rows(min_row=2, max_row=2, max_col=6, values_only=True)))
        if headers != ['No', 'Nama Barang', 'Kode Barang', 'Stok Barang', 'Harga Jual', 'Harga Beli']:
            raise ValueError(f'Unexpected inventory columns: {headers}')
        rows = []
        seen = set()
        for number, values in enumerate(sheet.iter_rows(min_row=3, max_col=6, values_only=True), 3):
            _, name, sku, stock, sell, cost = values
            if not name and not sku:
                continue
            if not sku or normalized(sku) in seen:
                raise ValueError(f'Missing or duplicate product at row {number}')
            if not isinstance(stock, (int, float)) or stock < 0 or int(stock) != stock:
                raise ValueError(f'Invalid stock at row {number}')
            row = dict(name=str(name or '').strip(), sku=str(sku).strip(), stock=int(stock),
                       selling_price=float(sell), cost_price=float(cost), source_row=number)
            Product(type='barang', **row)
            seen.add(normalized(sku))
            rows.append(row)
        if not rows:
            raise ValueError('Empty master is not allowed')
        return rows
    finally:
        workbook.close()


def plan_inventory(rows, existing, reference, digest):
    products, used = [], set()
    incoming_skus = {normalized(row['sku']) for row in rows}
    for row in rows:
        prior = unique_match(row, existing)
        if prior and (prior['id'] in used or (normalized(prior.get('sku')) != normalized(row['sku']) and normalized(prior.get('sku')) in incoming_skus)):
            prior = None
        tax_reference = unique_match(row, reference)
        product_id = prior['id'] if prior else str(uuid.uuid5(uuid.NAMESPACE_URL, 'latest-inventory:' + normalized(row['sku'])))
        if product_id in used:
            raise ValueError('Two new products resolve to the same existing product')
        used.add(product_id)
        attrs = tax_reference or {}
        doc = Product(**{**(prior or {}), **row, 'id': product_id,
                         'type': attrs.get('type', 'barang'),
                         'owner': attrs.get('owner'), 'tax': attrs.get('tax'),
                         'category': attrs.get('category', 'TIRE'),
                         'brand': attrs.get('brand', ''), 'size': attrs.get('size', ''),
                         'service_fee': attrs.get('service_fee', 0)}).model_dump()
        doc.update(active=True, import_source='inventory_20260930', source_digest=digest,
                   source_row=row['source_row'], source_sheet='Laporan Data Barang',
                   tax_reference_sheet='DATABASE' if tax_reference else None,
                   tax_reference_row=attrs.get('row'),
                   mapping_status='matched' if tax_reference else 'unmapped')
        products.append(doc)
    return products, [p['id'] for p in existing if p['id'] not in used]


async def replace(path, output):
    # The API holds the same OS lock. Never modify a running application's data.
    with single_writer(db.name):
        async with financial_lock:
            before = await snapshot(db)
            rows = read_inventory(path)
            digest = hashlib.sha256(path.read_bytes()).hexdigest()
            products, archive = plan_inventory(rows, before.get('products', []), dataset('store_current')['products'], digest)
            output.mkdir(parents=True, exist_ok=False)
            (output / 'backup.bson.json').write_text(json_util.dumps(before), encoding='utf-8')
            summary = dict(database=db.name, source=path.name, digest=digest, active=len(products),
                           matched=sum(p['mapping_status'] == 'matched' for p in products),
                           unmapped=sum(p['mapping_status'] == 'unmapped' for p in products), archived=len(archive))
            (output / 'plan.json').write_text(json.dumps({'summary': summary, 'products': products}, ensure_ascii=False, indent=2), encoding='utf-8')
            existing_ids = {p['id'] for p in before.get('products', [])}
            async with FinancialOperation('replace_inventory_20260930') as op:
                for product in products:
                    if product['id'] in existing_ids:
                        await op.update('products', product['id'], {'$set': product})
                    else:
                        await op.insert('products', product)
                for product_id in archive:
                    await op.update('products', product_id, {'$set': {'active': False}})
                actual = await db.products.find({'active': {'$ne': False}}, {'_id': 0}).to_list(None)
                assert len(actual) == len(products)
                actual_by_id = {p['id']: p for p in actual}
                for product in products:
                    assert all(actual_by_id[product['id']].get(k) == v for k, v in product.items())
                after = await snapshot(db)
                excluded = {'products', 'financial_operations'}
                assert fingerprint({k:v for k,v in before.items() if k not in excluded}) == fingerprint({k:v for k,v in after.items() if k not in excluded}), 'Historical data changed'
            summary['verified'] = True
            (output / 'result.json').write_text(json.dumps(summary, indent=2), encoding='utf-8')
            print(json.dumps(summary))


if __name__ == '__main__':
    stamp = datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S')
    asyncio.run(replace(ROOT / 'imports/workbooks/Laporan_Data_Barang_30-09-2026.xlsx',
                        ROOT / 'test_reports/spreadsheet' / ('inventory-replacement-' + stamp)))
